
from planilha import selecionar_planilha
from openpyxl import load_workbook

print("ETAPA 1: Selecionando a planilha...")

caminho = selecionar_planilha()
print("Arquivo selecionado com sucesso.")

print("\nETAPA 2: Verificando a aba Exclusões...")

workbook = load_workbook(
    caminho,
    read_only=True,
    data_only=True
)

if "Exclusões" not in workbook.sheetnames:
    workbook.close()
    raise Exception(
        "A aba 'Exclusões' não foi encontrada."
    )

aba = workbook["Exclusões"]

print("Aba encontrada:", aba.title)
print("Quantidade de linhas utilizadas:", aba.max_row)

print("\nCabeçalhos encontrados na linha 5:")

for coluna in range(1, aba.max_column + 1):
    valor = aba.cell(row=5, column=coluna).value

    if valor is not None:
        print(f"Coluna {coluna}: {valor}")

workbook.close()

print("\nTeste concluído. Nenhum dado foi alterado.")
